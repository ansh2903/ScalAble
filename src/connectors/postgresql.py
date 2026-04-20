import psycopg2
from psycopg2 import OperationalError
from sshtunnel import SSHTunnelForwarder
import duckdb
import pyarrow.parquet as pq
import urllib.parse
import pandas as pd
import openpyxl
import csv
import os
import io

from src.core.exception import CustomException
from src.core.logger import logging
from src.core.utils import is_running_in_docker


def _get_direct_conn(config):
    """Direct/cloud connection, with optional SSL."""
    return psycopg2.connect(
        host=config['host'],
        port=int(config['port']),
        user=config['user'],
        password=config['password'],
        dbname=config['dbname'],
        sslmode=config.get('sslmode', None),
        sslrootcert=config.get('sslrootcert', None),
        sslcert=config.get('sslcert', None),
        sslkey=config.get('sslkey', None),
        connect_timeout=15
    )


def _get_local_conn(config):
    """Local machine connection."""
    host = 'host.docker.internal' if is_running_in_docker() else '127.0.0.1'
    return psycopg2.connect(
        host=host,
        port=int(config['port']),
        user=config['user'],
        password=config['password'],
        dbname=config['dbname'],
        connect_timeout=5
    )


def _get_ssh_conn(config):
    """
    SSH tunnel connection.
    Returns (tunnel, conn) — caller must close both.
    The tunnel must stay open while conn is in use,
    so we return it for the caller to manage.
    """
    tunnel = SSHTunnelForwarder(
        (config['ssh_host'], int(config.get('ssh_port', 22))),
        ssh_username=config['ssh_user'],
        ssh_pkey=config.get('ssh_private_key'),
        remote_bind_address=(config['host'], int(config['port']))
    )
    tunnel.start()
    conn = psycopg2.connect(
        host='127.0.0.1',
        port=tunnel.local_bind_port,
        user=config['user'],
        password=config['password'],
        dbname=config['dbname'],
        connect_timeout=15
    )
    return tunnel, conn


def _get_conn(config):
    """
    Single entry point — returns (tunnel_or_None, conn).
    Callers always get a tuple so cleanup is consistent.
    """
    strategy = config.get('strategy')
    if strategy == 'local':
        return None, _get_local_conn(config)
    elif strategy == 'ssh':
        return _get_ssh_conn(config)
    else:
        return None, _get_direct_conn(config)


def _close(tunnel, conn, cursor=None):
    """Safe cleanup regardless of strategy."""
    try:
        if cursor:
            cursor.close()
    except Exception:
        pass
    try:
        if conn:
            conn.close()
    except Exception:
        pass
    try:
        if tunnel:
            tunnel.stop()
    except Exception:
        pass


# ── Public API ────────────────────────────────────────────────

def test_connection(config):
    tunnel, conn = None, None
    try:
        tunnel, conn = _get_conn(config)

        cursor = conn.cursor()

        # Set schema search path if provided
        schema = config.get('schema')
        if schema:
            cursor.execute(f"SET search_path TO {schema};")

        cursor.execute("SELECT 1;")
        _close(tunnel, conn, cursor)
        return True, "Connection successful"

    except OperationalError as e:
        _close(tunnel, conn)
        return False, str(e)
    except Exception as e:
        _close(tunnel, conn)
        return False, str(e)


def metadata(config):
    tunnel, conn = None, None
    try:
        tunnel, conn = _get_conn(config)
        cursor = conn.cursor()

        # Use provided schema, fall back to 'public'
        schema = config.get('schema') or 'public'

        db_meta = {
            "db_type": "postgresql",
            "database": config['dbname'],
            "schema": schema,
            "table_count": 0,
            "total_columns": 0,
            "tables": {}
        }

        # All base tables in the target schema
        cursor.execute("""
            SELECT table_name 
            FROM information_schema.tables 
            WHERE table_schema = %s AND table_type = 'BASE TABLE';
        """, (schema,))
        tables = cursor.fetchall()
        db_meta["table_count"] = len(tables)

        for (table_name,) in tables:
            # Columns
            cursor.execute("""
                SELECT column_name, data_type 
                FROM information_schema.columns 
                WHERE table_schema = %s AND table_name = %s;
            """, (schema, table_name))
            cols = cursor.fetchall()
            column_names = [c[0] for c in cols]
            column_types = {c[0]: c[1] for c in cols}
            db_meta["total_columns"] += len(column_names)

            # Row count
            cursor.execute(f'SELECT COUNT(*) FROM "{schema}"."{table_name}";')
            row_count = cursor.fetchone()[0]

            # Size
            qualified = f'"{schema}"."{table_name}"'
            cursor.execute("SELECT pg_total_relation_size(%s);", (f"{schema}.{table_name}",))
            size_bytes = cursor.fetchone()[0]
            cursor.execute("SELECT pg_size_pretty(pg_total_relation_size(%s));", (f"{schema}.{table_name}",))
            size_pretty = cursor.fetchone()[0]

            # Primary keys
            cursor.execute("""
                SELECT a.attname
                FROM pg_index i
                JOIN pg_attribute a ON a.attrelid = i.indrelid AND a.attnum = ANY(i.indkey)
                WHERE i.indrelid = %s::regclass AND i.indisprimary;
            """, (f"{schema}.{table_name}",))
            pk_columns = [r[0] for r in cursor.fetchall()]

            # PK bounds
            first_pk, last_pk = None, None
            if pk_columns:
                pk = pk_columns[0]
                cursor.execute(f'SELECT "{pk}" FROM "{schema}"."{table_name}" ORDER BY "{pk}" ASC LIMIT 1;')
                r = cursor.fetchone()
                first_pk = r[0] if r else None
                cursor.execute(f'SELECT "{pk}" FROM "{schema}"."{table_name}" ORDER BY "{pk}" DESC LIMIT 1;')
                r = cursor.fetchone()
                last_pk = r[0] if r else None

            # Unique constraints
            cursor.execute("""
                SELECT a.attname
                FROM pg_constraint c
                JOIN pg_class t ON c.conrelid = t.oid
                JOIN pg_namespace n ON n.oid = t.relnamespace
                JOIN pg_attribute a ON a.attrelid = t.oid AND a.attnum = ANY(c.conkey)
                WHERE c.contype = 'u' AND t.relname = %s AND n.nspname = %s;
            """, (table_name, schema))
            unique_keys = [r[0] for r in cursor.fetchall()]

            # Foreign keys
            cursor.execute("""
                SELECT kcu.column_name,
                       ccu.table_name  AS ref_table,
                       ccu.column_name AS ref_column
                FROM information_schema.table_constraints tc
                JOIN information_schema.key_column_usage kcu
                  ON tc.constraint_name = kcu.constraint_name
                 AND tc.table_schema    = kcu.table_schema
                JOIN information_schema.constraint_column_usage ccu
                  ON ccu.constraint_name = tc.constraint_name
                WHERE tc.constraint_type = 'FOREIGN KEY'
                  AND tc.table_schema    = %s
                  AND tc.table_name      = %s;
            """, (schema, table_name))
            foreign_keys = [
                {"column": r[0], "references_table": r[1], "references_column": r[2]}
                for r in cursor.fetchall()
            ]

            db_meta["tables"][table_name] = {
                "columns":       column_names,
                "column_types":  column_types,
                "row_count":     row_count,
                "size_bytes":    size_bytes,
                "size_pretty":   size_pretty,
                "primary_keys":  pk_columns,
                "row_bounds":    {"first_pk": first_pk, "last_pk": last_pk},
                "unique_keys":   unique_keys,
                "foreign_keys":  foreign_keys,
                "candidate_keys": list(set(pk_columns + unique_keys)),
            }

        _close(tunnel, conn, cursor)
        return True, db_meta

    except OperationalError as e:
        _close(tunnel, conn)
        return False, str(e)
    except Exception as e:
        _close(tunnel, conn)
        return False, str(e)


def file_to_db(config, file_path, table_name, ext):
    """Load a file into an existing table. Uses same config dict as everything else."""
    tunnel, conn = None, None
    cur = None
    schema = config.get('schema') or 'public'
    qualified = f'"{schema}"."{table_name}"'

    try:
        tunnel, conn = _get_conn(config)
        cur = conn.cursor()

        if ext in (".csv", ".txt"):
            with open(file_path, "r", encoding="utf-8") as f:
                cur.copy_expert(f"COPY {qualified} FROM STDIN WITH CSV HEADER", f)

        elif ext in (".xls", ".xlsx"):
            wb = openpyxl.load_workbook(file_path, data_only=True)
            buf = io.StringIO()
            csv.writer(buf).writerows(wb.active.iter_rows(values_only=True))
            buf.seek(0)
            cur.copy_expert(f"COPY {qualified} FROM STDIN WITH CSV HEADER", buf)
            buf.close()

        elif ext == ".json":
            try:
                df = pd.read_json(file_path, lines=True)
            except ValueError:
                df = pd.read_json(file_path)
            df.columns = [str(c).strip().replace(" ", "_") for c in df.columns]
            buf = io.StringIO()
            df.to_csv(buf, index=False)
            buf.seek(0)
            cur.copy_expert(f"COPY {qualified} FROM STDIN WITH CSV HEADER", buf)
            buf.close()

        else:
            return {"ok": False, "error": f"Unsupported file type: {ext}"}

        conn.commit()
        return {"ok": True}

    except Exception as e:
        if conn:
            conn.rollback()
        return {"ok": False, "error": str(e)}
    finally:
        _close(tunnel, conn, cur)

# For iterator generation
def data_extraction(stream_name, query, config, memory_ceiling=None, batch_row_size=None):
    tunnel = None
    conn = None
    writer = None
    stream_dir = f'src/temp/{stream_name}'
    os.makedirs(stream_dir, exist_ok=True)

    source_path = f'{stream_dir}/source.parquet'
    working_path = f'{stream_dir}/working.parquet'

    try:
        tunnel, duckdb_uri = _setup_duckdb_route(config)
        conn = duckdb.connect(':memory:')

        conn.execute(f"SET memory_limit = '{memory_ceiling if memory_ceiling else '1GB'}';")
        conn.execute("INSTALL postgres; LOAD postgres;")
        conn.execute(f"ATTACH '{duckdb_uri}' AS remote_db (TYPE POSTGRES);")
        conn.execute('USE remote_db;')

        schema = config.get('schema')
        if schema:
            conn.execute(f"SET search_path = {schema};")

        batch_size = batch_row_size if batch_row_size else 10
        reader = conn.sql(query).fetch_arrow_reader(batch_size)
        logging.info('reader object for data extraction created')
        
        fmt = 'parquet'
        if fmt == 'parquet':
            for batch in reader:
                if writer is None:
                    writer = pq.ParquetWriter(source_path, batch.schema, compression='snappy')

                writer.write_batch(batch)

        else:
            raise Exception

        logging.info(f'directory {stream_dir} created, data added')

        return 'success', stream_dir
    
    except Exception as e:
        logging.info(f"Streaming failed: {str(e)}")
        return {'status': 'error', 'message': str(e)}
    
    finally:
        if writer: writer.close()
        if tunnel: tunnel.close()
        if conn: conn.close()

# For preview of the view
def preview_execute(query, config):
    tunnel = None
    conn = None
    try:
        tunnel, duckdb_uri = _setup_duckdb_route(config)
        conn = duckdb.connect(':memory:')
        conn.execute("SET memory_limit = '1GB';") 
        conn.execute("INSTALL postgres; LOAD postgres;")
        conn.execute(f"ATTACH '{duckdb_uri}' AS remote_db (TYPE POSTGRES);")
        conn.execute('USE remote_db;')

        schema = config.get('schema')
        if schema:
            conn.execute(f"SET search_path = {schema};")

        rel = conn.sql(query)
        yield {"type": "columns", "content": rel.columns}

        try:
            count_query = f"SELECT COUNT(*) FROM ({str(query).strip(';')}) AS _count_tbl;"
            total_count = conn.execute(count_query).fetchone()[0]
        except Exception:
            total_count = "Unknown"

        preview_data = rel.limit(100).fetch_arrow_table().to_pylist()

        yield {
            "type": "metadata",
            "total_rows": total_count,
            "preview_count": len(preview_data)
        }

        yield {"type": "rows", "content": preview_data}

    except Exception as e:
        yield {"type": "error", "content": str(e)}
    finally:
        if conn: conn.close()
        if tunnel: tunnel.stop()

def _setup_duckdb_route(config):
    tunnel = None
    strategy = config.get('strategy')

    target_host = config.get('host', '127.0.0.1')
    target_port = config['port']

    if strategy == 'local':
        target_host = 'host.docker.internal' if is_running_in_docker() else '127.0.0.1'

    elif strategy == 'ssh':
        tunnel = SSHTunnelForwarder(
            (config['ssh_host'], int(config.get('ssh_port', 22))),
            ssh_username=config['ssh_user'],
            ssh_pkey=config.get('ssh_private_key'),
            remote_bind_address=(config['host'], int(config['port']))
        )

        tunnel.start()
        target_host = '127.0.0.1'
        target_port = tunnel.local_bind_port

    user = urllib.parse.quote_plus(config['user'])
    password = urllib.parse.quote_plus(config['password'])
    dbname = urllib.parse.quote_plus(config['dbname'])

    uri = f"postgresql://{user}:{password}@{target_host}:{target_port}/{dbname}"

    if strategy not in ['local', 'ssh']:
        ssl_params = []
        for key in ['sslmode', 'sslrootcert', 'sslcert', 'sslkey']:
            if config.get(key):
                ssl_params.append(f"{key}={urllib.parse.quote_plus(config[key])}")
        
        if ssl_params:
            uri += "?" + "&".join(ssl_params)

    return tunnel, uri
    

# arrow flight for later implementation 
'''
import pyarrow.flight as fl

class StreamFlightServer(fl.FlightServerBase):
    def __init__(self):
        super().__init__("grpc://localhost:8815")
        self._streams = {}  # stream_id → RecordBatchReader

    def register(self, stream_id, reader):
        self._streams[stream_id] = reader

    def do_get(self, context, ticket):
        stream_id = ticket.ticket.decode()
        reader = self._streams.get(stream_id)
        if not reader:
            raise fl.FlightServerError("Stream not found")
        return fl.RecordBatchStream(reader)

import pyarrow.flight as fl
client = fl.connect("grpc://localhost:8815")
reader = client.do_get(fl.Ticket(b"your-stream-id"))
batch = reader.read_chunk().data  # Arrow RecordBatch
'''